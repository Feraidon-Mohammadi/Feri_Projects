# pip install python-docx
# install that file is important


from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

########################################################################################################################
print(f"\n")
title = " first file ".upper()
print(title.center(80, "="))
print()
########################################################################################################################






# Create a new Word document
doc = Document()

# Add a title
title = doc.add_heading('My Word Document', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add text with different formatting
p = doc.add_paragraph('This is a sample text.')
p.add_run(' Bold Text').bold = True
p.add_run(' Italic Text').italic = True

# Change font size and style
run = doc.add_paragraph('Different font size and style:')
run.add_run(' 12pt Bold Red Text').bold = True
run.add_run(' 10pt Italic Blue Text').italic = True

for run in run.runs:
    run.font.size = Pt(12) if 'Bold' in run.text else Pt(10)
    run.font.color.rgb = (0, 0, 255) if 'Blue' in run.text else (255, 0, 0)

# Add a date-time
now = datetime.now()
current_time = now.strftime("%Y-%m-%d %H:%M:%S")
doc.add_paragraph(f'Current Date and Time: {current_time}')

# Save the document
doc.save('sample.docx')